import json
import os
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from urllib.parse import urlparse, urlunparse
import yaml

from reports.report_sanitizer import sanitize_report_tree

USER_PASS_LIST = [
    "Usability", "Security", "Efficiency", "Reliability",
    "Performance", "Accessibility", "Scalability", "Stability"
]
STATUS_ORDER = ["FAIL", "PASS", "SKIPPED"]
SEVERITY_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO", "N/A"]


def _normalize_target_path(target_url, path):
    target = str(target_url or "").strip().rstrip("/")
    raw_path = str(path or "").strip()

    if not target:
        return raw_path or "N/A"
    if not raw_path:
        return target

    target_parts = urlparse(target)
    if raw_path.startswith(("http://", "https://")):
        parsed = urlparse(raw_path)
        if parsed.netloc and parsed.netloc != target_parts.netloc:
            return urlunparse(
                (
                    target_parts.scheme or parsed.scheme,
                    target_parts.netloc,
                    parsed.path or "/",
                    parsed.params,
                    parsed.query,
                    parsed.fragment,
                )
            )
        return raw_path

    methodless_path = raw_path
    if " " in raw_path and raw_path.split(" ", 1)[0].isalpha():
        _, maybe_path = raw_path.split(" ", 1)
        if maybe_path.startswith("/"):
            methodless_path = maybe_path

    if methodless_path.startswith("/"):
        return f"{target}{methodless_path}"

    return raw_path


def _first_list(*candidates):
    for candidate in candidates:
        if isinstance(candidate, list) and candidate:
            return candidate
    return []


class ExcelReportGenerator:
    def __init__(self, output_dir="reports/output"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def _load_traffic_inventory(self, raw_results):
        metadata = raw_results.get("metadata", {}) if isinstance(raw_results, dict) else {}
        traffic_capture = metadata.get("traffic_capture", {}) if isinstance(metadata, dict) else {}
        if not isinstance(traffic_capture, dict):
            traffic_capture = {}

        if traffic_capture.get("captured_requests"):
            return traffic_capture

        candidate_paths = [
            traffic_capture.get("traffic_inventory_json_path"),
            traffic_capture.get("traffic_inventory_path"),
            os.path.join("keploy", "traffic", "traffic-inventory.json"),
            os.path.join("keploy", "tests", "keploy", "traffic-inventory.yaml"),
        ]

        for candidate in candidate_paths:
            if not candidate:
                continue
            path = Path(candidate)
            if not path.exists():
                continue

            try:
                with open(path, "r", encoding="utf-8") as handle:
                    if path.suffix.lower() == ".json":
                        loaded = json.load(handle)
                    else:
                        loaded = yaml.safe_load(handle)
            except Exception:
                continue

            if isinstance(loaded, dict):
                if loaded.get("captured_requests"):
                    captured_requests = loaded.get("captured_requests", [])
                    return {
                        "captured_requests": captured_requests,
                        "captured_request_count": len(captured_requests),
                        "unique_endpoint_count": loaded.get("unique_api_count", 0),
                        "unique_hosts": sorted({req.get("host") for req in captured_requests if isinstance(req, dict) and req.get("host")}),
                        "traffic_inventory_path": str(path),
                        "traffic_inventory_json_path": str(path),
                    }

                capture_profile = loaded.get("capture_profile", {})
                if isinstance(capture_profile, dict) and capture_profile.get("captured_requests"):
                    return capture_profile

        return traffic_capture

    def _resolve_discovery_baseline(self, raw_results):
        metadata = raw_results.get("metadata", {}) if isinstance(raw_results, dict) else {}
        service_mesh = metadata.get("service_mesh_discovery", {}) if isinstance(metadata, dict) else {}
        scenarios = raw_results.get("scenarios", []) if isinstance(raw_results, dict) else []

        discovery_baseline = _first_list(
            metadata.get("discovery_baseline"),
            service_mesh.get("baseline") if isinstance(service_mesh, dict) else None,
            raw_results.get("discovered_apis"),
            raw_results.get("baseline"),
            raw_results.get("discovery_baseline"),
            metadata.get("discovered_apis"),
            metadata.get("inventory_summary", {}).get("discovered_apis") if isinstance(metadata.get("inventory_summary"), dict) else None,
        )
        if discovery_baseline:
            return discovery_baseline

        derived = []
        seen = set()
        if isinstance(scenarios, list):
            for scenario in scenarios:
                if not isinstance(scenario, dict):
                    continue
                for endpoint in scenario.get("endpoints", []) if isinstance(scenario.get("endpoints", []), list) else []:
                    if not isinstance(endpoint, dict):
                        continue
                    method = str(endpoint.get("method") or "GET").upper()
                    raw_path = str(endpoint.get("path") or "").strip()
                    raw_url = str(endpoint.get("url") or "").strip()
                    path = raw_path
                    if not path and raw_url:
                        parsed = urlparse(raw_url)
                        path = parsed.path or raw_url
                    if not path:
                        continue
                    key = (method, path)
                    if key in seen:
                        continue
                    seen.add(key)
                    derived.append({
                        "method": method,
                        "path": path,
                        "status": endpoint.get("status", 200),
                        "description": endpoint.get("expected_behavior") or endpoint.get("description") or "Discovered Interface",
                        "source": endpoint.get("source", "scenario"),
                    })

        return derived

    def generate(self, raw_results, scorecard, final_score, tier, filepath="user_pass_certification_report.xlsx", report_insights=None):
        print("[*] Generating Professionally Styled Excel Report...")
        target_url = raw_results.get("metadata", {}).get("target_url", "")
        raw_results = sanitize_report_tree(raw_results, target_url)
        traffic_capture = self._load_traffic_inventory(raw_results)
        out_path = os.path.join(self.output_dir, filepath)
        try:
            with open(out_path, "a", encoding="utf-8"):
                pass
        except PermissionError:
            fallback_name = f"{Path(filepath).stem}_clean{Path(filepath).suffix}"
            out_path = os.path.join(self.output_dir, fallback_name)
            print(f"[!] Report file is locked. Writing Excel output to {out_path}")

        summary_data = []
        for attr_key in USER_PASS_LIST:
            attr_lower = attr_key.lower()
            sc = scorecard.get(attr_lower, {})
            summary_data.append({
                "Attribute": attr_key,
                "Tools Passed": sc.get("passed", 0),
                "Tools Failed": sc.get("failed", 0),
                "Tools Skipped": sc.get("skipped", 0),
                "Total Tools": sc.get("total_tools", 0),
                "Blockers": sc.get("blockers", 0),
                "Attribute Score (%)": sc.get("attribute_score", 0),
                "Weight (%)": round(sc.get("weight", 0) * 100),
                "Weighted Score": sc.get("weighted_contribution", 0),
                "Tier": sc.get("tier", "N/A"),
                "Zero-Defect": sc.get("verdict", "N/A"),
            })

        total_passed = sum(r["Tools Passed"] for r in summary_data)
        total_failed = sum(r["Tools Failed"] for r in summary_data)
        total_skipped = sum(r["Tools Skipped"] for r in summary_data)
        total_tools = sum(r["Total Tools"] for r in summary_data)
        total_blockers = sum(r["Blockers"] for r in summary_data)

        summary_data.append({
            "Attribute": "FINAL SCORECARD",
            "Tools Passed": total_passed,
            "Tools Failed": total_failed,
            "Tools Skipped": total_skipped,
            "Total Tools": total_tools,
            "Blockers": total_blockers,
            "Attribute Score (%)": None,
            "Weight (%)": 100,
            "Weighted Score": final_score,
            "Tier": tier.upper(),
            "Zero-Defect": tier.upper(),
        })
        df_summary = pd.DataFrame(summary_data)

        findings_data = []
        for attr_lower, data in raw_results.items():
            if attr_lower == "metadata":
                continue
            attr_cap = attr_lower.capitalize()
            if isinstance(data, dict) and "findings" in data:
                for tool, details in data["findings"].items():
                    success = details.get("success", False)
                    status = details.get("status", "PASS" if success else "FAIL")
                    mode = details.get("mode", "LIVE")
                    findings_data.append({
                        "Attribute": attr_cap,
                        "Tool Name": tool.replace("_", " ").title(),
                        "Status": status,
                        "Execution Mode": mode,
                        "Severity": details.get("severity", "INFO" if success else "CRITICAL"),
                        "Target Path": _normalize_target_path(target_url, details.get("path", "-")),
                        "Diagnostic Log": details.get("log", ""),
                        "Recommendation": details.get("recommendation", "N/A"),
                    })

        if not findings_data:
            execution_notice = raw_results.get("metadata", {}).get("execution_reason", "No tools were executed.")
            findings_data.append({
                "Attribute": "Execution",
                "Tool Name": "No Tools Executed",
                "Status": raw_results.get("metadata", {}).get("execution_state", "SKIPPED"),
                "Execution Mode": "N/A",
                "Severity": "INFO",
                "Target Path": target_url or "N/A",
                "Diagnostic Log": execution_notice,
                "Recommendation": "Verify the strategy output includes a non-empty tools list or fix tool name mapping.",
            })

        findings_columns = [
            "Attribute", "Tool Name", "Status", "Execution Mode",
            "Severity", "Target Path", "Diagnostic Log", "Recommendation"
        ]
        df_findings = pd.DataFrame(findings_data, columns=findings_columns)
        if not df_findings.empty:
            df_findings["Attribute"] = pd.Categorical(df_findings["Attribute"], categories=USER_PASS_LIST, ordered=True)
            df_findings["Status"] = pd.Categorical(df_findings["Status"], categories=STATUS_ORDER, ordered=True)
            df_findings["Severity"] = pd.Categorical(df_findings["Severity"], categories=SEVERITY_ORDER, ordered=True)
            df_findings = df_findings.sort_values(by=["Attribute", "Status", "Severity"], ascending=[True, True, True])

        df_failures = df_findings[df_findings["Status"] == "FAIL"].copy() if not df_findings.empty else df_findings.copy()

        discovery_data = self._resolve_discovery_baseline(raw_results)
        if not discovery_data:
            discovery_data = [
                {
                    "Message": "No confirmed APIs discovered for this target. The capture contained page routes or non-API traffic only."
                }
            ]
        df_baseline = pd.DataFrame(discovery_data)

        active_resilience = report_insights or raw_results.get("metadata", {}).get("active_resilience", {})
        active_rows = []
        if isinstance(active_resilience, dict) and active_resilience:
            # Handle both raw nested structure (from scorer) and flat summary dict
            capture = active_resilience.get("capture_profile") or {}
            coverage = active_resilience.get("coverage_booster") or {}
            chaos = active_resilience.get("chaos_profile") or {}
            pii = active_resilience.get("pii_audit") or {}
            drift = active_resilience.get("behavioral_drift") or {}

            # If no sub-keys found, the dict is already a flat summary
            is_flat = not any([capture, coverage, chaos, pii, drift])

            if is_flat:
                # Use flat keys from _build_active_resilience_summary
                active_rows = [
                    {"Metric": "Capture Mode", "Value": active_resilience.get("capture_mode", "N/A")},
                    {"Metric": "OpenAPI Source", "Value": active_resilience.get("openapi_source") or "N/A"},
                    {"Metric": "OpenAPI Trace Count", "Value": active_resilience.get("openapi_trace_count", 0)},
                    {"Metric": "Chaos Mutated Files", "Value": active_resilience.get("chaos_mutated_files", 0)},
                    {"Metric": "Chaos Mutations", "Value": ", ".join(active_resilience.get("chaos_mutations", [])) or "N/A"},
                    {"Metric": "PII Files Scanned", "Value": active_resilience.get("pii_files_scanned", 0)},
                    {"Metric": "PII Leaks", "Value": active_resilience.get("pii_leaks", 0)},
                    {"Metric": "Behavioral Similarity (%)", "Value": active_resilience.get("drift_similarity_percent", 100.0)},
                    {"Metric": "Behavioral Drift (%)", "Value": active_resilience.get("drift_percent", 0.0)},
                    {"Metric": "Resilience Index (%)", "Value": active_resilience.get("resilience_index", 0.0)},
                    {"Metric": "Coverage Score (%)", "Value": active_resilience.get("coverage_percent", 0.0)},
                    {"Metric": "PII Score (%)", "Value": active_resilience.get("pii_score", 100.0)},
                    {"Metric": "Chaos Score (%)", "Value": active_resilience.get("chaos_score", 0.0)},
                ]
            else:
                # Use nested sub-key structure
                active_rows = [
                    {"Metric": "Capture Mode", "Value": capture.get("capture_mode", "N/A")},
                    {"Metric": "Capture Command", "Value": capture.get("command", "N/A")},
                    {"Metric": "Captured Requests", "Value": capture.get("captured_request_count", 0)},
                    {"Metric": "Unique Endpoints", "Value": capture.get("unique_endpoint_count", 0)},
                    {"Metric": "OpenAPI Source", "Value": coverage.get("source_path") or "N/A"},
                    {"Metric": "OpenAPI Trace Count", "Value": coverage.get("generated_trace_count", 0)},
                    {"Metric": "Chaos Mutated Files", "Value": len(chaos.get("mutated_files", [])) if isinstance(chaos, dict) else 0},
                    {"Metric": "Chaos Mutations", "Value": ", ".join(chaos.get("mutations_applied", [])) if isinstance(chaos, dict) and chaos.get("mutations_applied") else "N/A"},
                    {"Metric": "PII Files Scanned", "Value": pii.get("files_scanned", 0)},
                    {"Metric": "PII Leaks", "Value": len(pii.get("leaks_found", [])) if isinstance(pii, dict) else 0},
                    {"Metric": "Behavioral Similarity (%)", "Value": drift.get("behavioral_similarity_percent", 100.0)},
                    {"Metric": "Behavioral Drift (%)", "Value": drift.get("behavioral_drift_percent", 0.0)},
                    {"Metric": "Resilience Index (%)", "Value": active_resilience.get("resilience_index", 0.0)},
                    {"Metric": "Coverage Score (%)", "Value": active_resilience.get("coverage_percent", 0.0)},
                    {"Metric": "PII Score (%)", "Value": active_resilience.get("pii_score", 100.0)},
                    {"Metric": "Chaos Score (%)", "Value": active_resilience.get("chaos_score", 0.0)},
                ]
        df_active = pd.DataFrame(active_rows, columns=["Metric", "Value"])

        traffic_rows = []
        if isinstance(traffic_capture, dict):
            for request in traffic_capture.get("captured_requests", []):
                if not isinstance(request, dict):
                    continue
                traffic_rows.append({
                    "Method": request.get("method", "-"),
                    "Host": request.get("host", "-"),
                    "Path": request.get("path", "-"),
                    "Full URL": request.get("url", "-"),
                    "Status": request.get("status", "-"),
                    "Resource Type": request.get("resource_type", "-"),
                    "Navigation": request.get("is_navigation", False),
                })
        df_traffic = pd.DataFrame(traffic_rows, columns=["Method", "Host", "Path", "Full URL", "Status", "Resource Type", "Navigation"])

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df_summary.to_excel(writer, sheet_name="Scorecard Summary", index=False)
            df_baseline.to_excel(writer, sheet_name="Discovered API Baseline", index=False)
            df_active.to_excel(writer, sheet_name="Active Resilience", index=False)
            df_traffic.to_excel(writer, sheet_name="Traffic Inventory", index=False)
            df_findings.to_excel(writer, sheet_name="All Tool Findings", index=False)
            df_failures.to_excel(writer, sheet_name="Failures Requiring Action", index=False)

            workbook = writer.book
            ws_summary = workbook["Scorecard Summary"]
            self._apply_pro_styling(ws_summary, highlight_rows=[len(summary_data) + 1])
            self._apply_pro_styling(workbook["Discovered API Baseline"])
            self._apply_pro_styling(workbook["Active Resilience"])
            self._apply_pro_styling(workbook["Traffic Inventory"])
            self._apply_pro_findings_styling(workbook["All Tool Findings"])
            self._apply_pro_findings_styling(workbook["Failures Requiring Action"])

            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Attribute Scores (%)"
            chart.y_axis.title = "Score"
            chart.x_axis.title = "Attribute"
            data = Reference(ws_summary, min_col=7, min_row=1, max_row=9)
            cats = Reference(ws_summary, min_col=1, min_row=2, max_row=9)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws_summary.add_chart(chart, "K2")

        print(f"[+] Professionally styled Excel report saved -> {out_path}")
        return out_path

    def _apply_pro_styling(self, ws, highlight_rows=None):
        highlight_rows = highlight_rows or []
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        pass_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
        fail_fill = PatternFill(start_color="F1C1C1", end_color="F1C1C1", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = thin_border
                if cell.column == ws.max_column:
                    if cell.value == "PASS":
                        cell.fill = pass_fill
                    elif cell.value == "FAIL":
                        cell.fill = fail_fill

            if row[0].row in highlight_rows:
                for cell in row:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 20

        if ws.dimensions:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        if ws.title == "Scorecard Summary":
            ws.sheet_view.showGridLines = False

    def _apply_pro_findings_styling(self, ws):
        header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border

        col_widths = [18, 25, 12, 14, 15, 30, 50, 50]
        for i, width in enumerate(col_widths):
            ws.column_dimensions[chr(65 + i)].width = width

        for row in ws.iter_rows(min_row=2):
            status_cell = row[2] if ws.max_column > 2 else None
            mode_cell = row[3] if ws.max_column > 3 else None
            sev_cell = row[4] if ws.max_column > 4 else None

            if status_cell and status_cell.value == "FAIL":
                status_cell.font = Font(color="C00000", bold=True)
            elif status_cell and status_cell.value == "PASS":
                status_cell.font = Font(color="008000")
            elif status_cell and status_cell.value == "SKIPPED":
                status_cell.font = Font(color="808080", italic=True)

            if mode_cell and mode_cell.value == "SIMULATED":
                mode_cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
            elif mode_cell and mode_cell.value == "SKIPPED":
                mode_cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")

            if sev_cell:
                if sev_cell.value in ["CRITICAL", "HIGH"]:
                    sev_cell.fill = PatternFill(start_color="FFCCCC", fill_type="solid")
                elif sev_cell.value == "MEDIUM":
                    sev_cell.fill = PatternFill(start_color="FFF2CC", fill_type="solid")

            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        if ws.dimensions:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
