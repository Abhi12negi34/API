import json
from pathlib import Path

from openpyxl import load_workbook

from core.api.keploy_intercept_proxy import start_keploy_record
from core.api.playwright_stimulator import _attempt_browser_login, _build_api_probe_plan, _looks_like_login_url, _navigate_page, _probe_api_endpoints, _resolve_auth_credentials, _resolve_auth_settings, _serialize_auth_settings, simulate_user_flows
from core.api.discovery_inventory import build_discovery_inventory
from core.api_testing_engine import ApiTestingEngine, _normalize_target_path
from agents.api_execution_agent import execute_scenarios_core
from reports.generators.excel_generator import ExcelReportGenerator
from reports.user_pass_certification_scorer import CertificationScorer
from tools.reliability.keploy_coverage_booster import generate_openapi_traces
from tools.reliability.keploy_drift_certifier import calculate_behavioral_drift
from tools.security.keploy_pii_leak_auditor import audit_pii_data_leaks
from tools.stability.keploy_chaos_fuzzer import inject_chaos


def test_start_keploy_record_writes_artifacts(tmp_path):
    seed_dir = tmp_path / "keploy" / "tests" / "keploy"
    seed_dir.mkdir(parents=True)
    seed_file = seed_dir / "discovery-baseline.yaml"
    seed_file.write_text(
        """
discovered_apis:
  - method: GET
    path: /api/v1/users
    status: 200
    description: Users list
  - method: POST
    path: /api/v1/users
    status: 201
    description: Create user
""".strip(),
        encoding="utf-8",
    )

    profile = start_keploy_record("https://example.com", workspace_root=tmp_path)

    baseline_path = Path(profile["baseline_path"])
    capture_path = tmp_path / "keploy" / "capture-profile.json"

    assert len(profile["baseline"]) >= 2
    assert any(api["path"] == "/api/v1/users" for api in profile["baseline"])
    assert profile["discovery_source_files"]
    assert baseline_path.exists()
    assert capture_path.exists()


def test_start_keploy_record_reads_traffic_inventory(tmp_path):
    seed_dir = tmp_path / "keploy" / "tests" / "keploy"
    seed_dir.mkdir(parents=True)
    traffic_file = seed_dir / "traffic-inventory.yaml"
    traffic_file.write_text(
        """
captured_requests:
  - method: GET
    url: https://example.com/api/v1/users
    path: /api/v1/users
    host: example.com
    status: 200
    resource_type: xhr
""".strip(),
        encoding="utf-8",
    )

    profile = start_keploy_record("https://example.com", workspace_root=tmp_path)

    assert profile["baseline"]
    assert profile["baseline"][0]["path"] == "/api/v1/users"
    assert Path(profile["baseline_path"]).exists()


def test_build_api_probe_plan_uses_discovered_endpoints(monkeypatch):
    monkeypatch.setattr(
        "core.api.playwright_stimulator.build_discovery_inventory",
        lambda target_url, workspace_root=None: {
            "discovered_apis": [
                {"method": "GET", "path": "/api/v1/users", "source_file": "seed"},
                {"method": "POST", "path": "/api/v1/orders", "source_file": "seed"},
                {"method": "GET", "path": "/", "source_file": "seed"},
            ]
        },
    )

    plan = _build_api_probe_plan("https://example.com", workspace_root=Path.cwd(), max_probes=10)

    assert len(plan) == 2
    assert plan[0]["url"] == "https://example.com/api/v1/users"
    assert plan[1]["url"] == "https://example.com/api/v1/orders"


def test_probe_api_endpoints_records_fetches():
    class FakePage:
        def __init__(self):
            self.evaluated = []
            self.timeouts = []

        def evaluate(self, script, params):
            self.evaluated.append(params["url"])
            return {"ok": True, "status": 200}

        def wait_for_timeout(self, ms):
            self.timeouts.append(ms)

    page = FakePage()
    probes = [
        {"method": "GET", "url": "https://example.com/api/v1/users", "path": "/api/v1/users"},
        {"method": "GET", "url": "https://example.com/api/v1/orders", "path": "/api/v1/orders"},
    ]

    executed = _probe_api_endpoints(page, probes, per_probe_wait_ms=25)

    assert len(executed) == 2
    assert page.evaluated == [
        "https://example.com/api/v1/users",
        "https://example.com/api/v1/orders",
    ]
    assert page.timeouts == [25, 25]


def test_resolve_auth_settings_builds_storage_state_path(tmp_path, monkeypatch):
    monkeypatch.setenv("PLAYWRIGHT_EMAIL", "tester@example.com")
    monkeypatch.setenv("PLAYWRIGHT_PASSWORD", "secret-password")

    settings = _resolve_auth_settings({}, workspace_root=tmp_path)
    serialized = _serialize_auth_settings(settings)

    assert settings["email_env"] == "PLAYWRIGHT_EMAIL"
    assert settings["password_env"] == "PLAYWRIGHT_PASSWORD"
    assert Path(settings["storage_state_path"]).name == "playwright-storage-state.json"
    assert serialized["storage_state_path"].endswith("playwright-storage-state.json")


def test_resolve_auth_credentials_accepts_literal_pasted_values(monkeypatch):
    monkeypatch.delenv("PLAYWRIGHT_USERNAME", raising=False)
    monkeypatch.delenv("PLAYWRIGHT_EMAIL", raising=False)
    monkeypatch.delenv("PLAYWRIGHT_PASSWORD", raising=False)

    username, password = _resolve_auth_credentials(
        {
            "username": "",
            "email": "",
            "password": "",
            "username_env": "PLAYWRIGHT_USERNAME",
            "email_env": "tester@example.com",
            "password_env": "secret-password",
        }
    )

    assert username == "tester@example.com"
    assert password == "secret-password"


def test_forced_login_ignores_saved_storage_state(monkeypatch, tmp_path):
    calls = {"fills": [], "saved_state": None}

    class FakeLocator:
        def __init__(self, label):
            self.label = label

        def fill(self, value):
            calls["fills"].append((self.label, value))

        def click(self, timeout=None):
            return None

        def press(self, key):
            calls["fills"].append((self.label, key))

    class FakeContext:
        def storage_state(self, path):
            calls["saved_state"] = path

    class FakePage:
        url = "http://example.com/login"

        def __init__(self):
            self.context = FakeContext()

        def wait_for_load_state(self, *args, **kwargs):
            return None

        def wait_for_timeout(self, ms):
            return None

    monkeypatch.setattr("core.api.playwright_stimulator._state_file_exists", lambda path: True)
    monkeypatch.setattr(
        "core.api.playwright_stimulator._resolve_auth_credentials",
        lambda auth_settings: ("tester@example.com", "secret-password"),
    )
    monkeypatch.setattr("core.api.playwright_stimulator._page_login_signals", lambda page: True)
    monkeypatch.setattr(
        "core.api.playwright_stimulator._auto_detect_login_locators",
        lambda page: (FakeLocator("username"), FakeLocator("password"), FakeLocator("submit")),
    )

    result = _attempt_browser_login(
        FakePage(),
        {
            "reuse_storage_state": True,
            "storage_state_path": tmp_path / "playwright-storage-state.json",
        },
        target_url="http://example.com",
        force=True,
    )

    assert result["attempted"] is True
    assert result["success"] is True
    assert calls["fills"] == [
        ("username", "tester@example.com"),
        ("password", "secret-password"),
    ]
    assert calls["saved_state"].endswith("playwright-storage-state.json")


def test_looks_like_login_url_detects_common_login_paths():
    assert _looks_like_login_url("http://localhost:9000/app/login")
    assert _looks_like_login_url("https://example.com/signin")
    assert _looks_like_login_url("https://example.com/account/login")
    assert not _looks_like_login_url("https://example.com/dashboard")


def test_navigate_page_detects_blank_navigation():
    class FakePage:
        def __init__(self):
            self.url = "about:blank"

        def goto(self, url, wait_until=None, timeout=None):
            return None

        def wait_for_load_state(self, state, timeout=None):
            return None

    assert _navigate_page(FakePage(), "http://example.com") is False


def test_simulate_user_flows_restores_start_url():
    class FakeMouse:
        def wheel(self, x, y):
            return None

    class FakeLink:
        def __init__(self, page):
            self.page = page

        def click(self, timeout=None):
            self.page.url = "about:blank"

    class FakeLocator:
        def __init__(self, page, selector):
            self.page = page
            self.selector = selector

        def count(self):
            return 0 if self.selector == "input[type='text']" else 1

        @property
        def first(self):
            return self

        def nth(self, index):
            return FakeLink(self.page)

        def fill(self, value):
            return None

        def press(self, value):
            return None

    class FakePage:
        def __init__(self):
            self.url = "http://example.com/app/orders"
            self.mouse = FakeMouse()
            self.goto_calls = []

        def locator(self, selector):
            return FakeLocator(self, selector)

        def wait_for_timeout(self, ms):
            return None

        def goto(self, url, wait_until=None, timeout=None):
            self.goto_calls.append(url)
            self.url = url

        def go_back(self):
            self.url = "about:blank"

    page = FakePage()
    simulate_user_flows(page)

    assert page.url == "http://example.com/app/orders"
    assert page.goto_calls[-1] == "http://example.com/app/orders"


def test_start_keploy_record_filters_static_and_foreign_entries(tmp_path):
    seed_dir = tmp_path / "keploy" / "tests" / "keploy"
    seed_dir.mkdir(parents=True)
    seed_file = seed_dir / "mixed-discovery.yaml"
    seed_file.write_text(
        """
capture_profile:
  target_url: https://example.com
  discovered_apis:
    - method: GET
      path: /api/v1/users
      status: 200
      description: Users list
    - method: GET
      path: /_next/static/chunks/app.js
      status: 200
      description: Static asset
    - method: GET
      path: /api/v1/orders
      status: 200
      target_url: https://foreign.example.com
      description: Foreign host request
""".strip(),
        encoding="utf-8",
    )

    profile = start_keploy_record("https://example.com", workspace_root=tmp_path)

    assert any(api["path"] == "/api/v1/users" for api in profile["baseline"])
    assert all(api["path"] != "/_next/static/chunks/app.js" for api in profile["baseline"])
    assert all(api.get("observed_host", "") != "foreign.example.com" for api in profile["baseline"])
    assert profile["discovered_api_count"] >= 1
    assert profile["inventory_summary"]["excluded_request_count"] == 2
    assert profile["inventory_summary"]["foreign_host_request_count"] == 1


def test_normalize_target_path_swaps_host():
    normalized = _normalize_target_path("https://www.myntra.com", "https://www.w3schools.com/api/v1/auth")
    assert normalized == "https://www.myntra.com/api/v1/auth"


def test_generate_openapi_traces_from_schema(tmp_path):
    spec_path = tmp_path / "openapi.yaml"
    spec_path.write_text(
        """
openapi: 3.0.0
paths:
  /api/v1/users:
    get: {}
  /api/v1/orders:
    post: {}
""".strip(),
        encoding="utf-8",
    )

    result = generate_openapi_traces("https://example.com", workspace_root=tmp_path)

    assert result["success"] is True
    assert result["generated_trace_count"] == 2
    assert Path(result["path"]).exists()


def test_inject_chaos_mutates_mocks(tmp_path):
    mocks_dir = tmp_path / "keploy" / "mocks"
    mocks_dir.mkdir(parents=True)
    (mocks_dir / "mocks.yaml").write_text(
        """
mocks:
  - name: payment_gateway
    response:
      status_code: 200
      body:
        ok: true
""".strip(),
        encoding="utf-8",
    )

    result = inject_chaos(mock_path=mocks_dir, workspace_root=tmp_path)

    assert result["success"] is True
    mutated_path = Path(result["mutated_files"][0])
    assert mutated_path.exists()
    mutated_text = mutated_path.read_text(encoding="utf-8")
    assert "503" in mutated_text or "delay_ms" in mutated_text


def test_pii_auditor_scans_keploy_artifacts(tmp_path):
    target_dir = tmp_path / "keploy" / "tests" / "test-set-1"
    target_dir.mkdir(parents=True)
    (target_dir / "recording.yaml").write_text(
        """
response:
  body:
    email: user@example.com
    api_key: supersecretapikeyvalue12345
""".strip(),
        encoding="utf-8",
    )

    result = audit_pii_data_leaks("https://example.com", workspace_root=tmp_path)

    assert result["success"] is False
    assert result["leaks_found"]


def test_behavioral_drift_decreases_on_change(tmp_path):
    report_path = tmp_path / "reports" / "output" / "user_pass_certification_report.json"
    report_path.parent.mkdir(parents=True)
    baseline_report = {
        "detailed_findings_by_attribute": {
            "reliability": {
                "findings": [
                    {"tool_id": "coverage_booster", "success": True, "status": "PASS", "severity": "INFO"},
                    {"tool_id": "chaos_stab", "success": True, "status": "PASS", "severity": "INFO"},
                ]
            }
        }
    }
    report_path.write_text(json.dumps(baseline_report), encoding="utf-8")

    current_results = {
        "metadata": {},
        "reliability": {
            "findings": {
                "coverage_booster": {"success": True, "status": "PASS", "severity": "INFO"},
                "chaos_stab": {"success": True, "status": "PASS", "severity": "INFO"},
            }
        },
    }
    result_same = calculate_behavioral_drift(current_results, workspace_root=tmp_path)
    assert result_same["behavioral_similarity_percent"] == 100.0

    current_results["reliability"]["findings"]["chaos_stab"] = {
        "success": False,
        "status": "FAIL",
        "severity": "CRITICAL",
    }
    result_changed = calculate_behavioral_drift(current_results, workspace_root=tmp_path)
    assert result_changed["behavioral_similarity_percent"] < 100.0


def test_build_analysis_brief_reports_failures_and_mapping_gaps():
    engine = ApiTestingEngine({"zaproxy": {}})
    raw_results = {
        "metadata": {
            "target_url": "https://example.com",
            "planning_context": "x" * 2000,
            "discovery_baseline": [{"method": "GET", "path": "/api/v1/users"}],
            "traffic_capture": {
                "capture_mode": "proxy",
                "captured_request_count": 1,
                "unique_endpoint_count": 1,
                "unique_hosts": ["example.com"],
                "discovery_source_files": ["seed.yaml"],
                "inventory_summary": {"discovered_api_count": 1},
            },
        },
        "security": {
            "overall": False,
            "findings": {
                "secure_config": {
                    "success": False,
                    "status": "FAIL",
                    "mode": "LIVE",
                    "severity": "HIGH",
                    "path": "https://example.com/api/v1/health",
                    "log": "A" * 500,
                    "recommendation": "B" * 400,
                }
            },
        },
    }

    brief = engine.build_analysis_brief(raw_results, max_failures=5)

    assert brief["discovery_summary"]["discovered_api_count"] == 1
    assert brief["failures"][0]["tool_id"] == "secure_config"
    assert brief["failures"][0]["log"].endswith("...")
    assert brief["mapping_gaps"]["configured_tool_count"] >= brief["mapping_gaps"]["implemented_tool_count"]
    assert len(brief["planning_context_excerpt"]) == 1200


def test_execute_scenarios_tool_emits_scorecard_shape():
    payload = {
        "target_url": "https://example.com",
        "run_all_registry_tools": False,
        "scenarios": [
            {
                "scenario_id": 1,
                "scenario_name": "Smoke",
                "endpoints": [
                    {
                        "method": "GET",
                        "url": "https://example.com/api/v1/health",
                        "source": "traffic-inventory",
                        "expected_behavior": "Returns OK",
                    }
                ],
                "tools": ["k6", "zap", "axe", "webvitals"],
            }
        ],
    }

    result = json.loads(execute_scenarios_core(json.dumps(payload)))

    assert result["metadata"]["total_scenarios"] == 1
    assert "performance" in result and "security" in result
    assert result["results"]["1_Smoke"]["k6"]["path"] == "https://example.com/api/v1/health"


def test_execute_scenarios_tool_accepts_list_payload():
    payload = [
        {
            "scenario_id": "product_catalog_browsing",
            "scenario_name": "Public Product Catalog Browsing",
            "endpoints": [
                {
                    "method": "GET",
                    "url": "https://example.com/api/v1/products/list",
                    "source": "traffic-inventory",
                    "expected_behavior": "Returns products",
                }
            ],
            "tools": ["k6", "bombardier"],
        }
    ]

    result = json.loads(execute_scenarios_core(json.dumps(payload)))

    assert result["metadata"]["total_scenarios"] == 1
    assert result["results"]["product_catalog_browsing_Public Product Catalog Browsing"]["k6"]["status"] == "PASS"


def test_grounded_strategy_steps_are_dynamic(tmp_path):
    from main import _build_grounded_strategy_payload

    get_payload = _build_grounded_strategy_payload(
        "https://example.com",
        {
            "discovered_apis": [
                {
                    "method": "GET",
                    "path": "/api/v1/orders",
                    "full_url": "https://example.com/api/v1/orders",
                    "description": "Orders list",
                }
            ]
        },
        None,
    )

    post_payload = _build_grounded_strategy_payload(
        "https://example.com",
        {
            "discovered_apis": [
                {
                    "method": "POST",
                    "path": "/api/v1/auth/session",
                    "full_url": "https://example.com/api/v1/auth/session",
                    "description": "Create session",
                }
            ]
        },
        None,
    )

    get_steps = get_payload["scenarios"][0]["steps"]
    post_steps = post_payload["scenarios"][0]["steps"]

    assert len(get_steps) == 4
    assert len(post_steps) >= 5
    assert post_steps[-1]["action"]


def test_shared_registry_includes_config_only_tools():
    engine = ApiTestingEngine({"zaproxy": {}})
    registry = engine._build_tool_registry("https://example.com")

    for tool_name in [
        "api_doc_clarity",
        "ui_visual_regress",
        "rtl_i18n_audit",
        "sql_injection_url",
        "cia_integrity",
        "rbac_access_audit",
        "cors_misconfig",
    ]:
        assert tool_name in registry


def test_scorer_normalizes_discovery_and_traffic_metadata():
    scorer = CertificationScorer(
        {
            "usability": 0.1,
            "security": 0.2,
            "efficiency": 0.05,
            "reliability": 0.15,
            "performance": 0.15,
            "accessibility": 0.1,
            "scalability": 0.15,
            "stability": 0.1,
        }
    )
    payload = {
        "metadata": {"target_url": "https://example.com"},
        "discovered_apis": [{"method": "GET", "path": "/api/v1/users"}],
        "traffic_requests": [
            {
                "method": "GET",
                "path": "/api/v1/users",
                "host": "example.com",
                "url": "https://example.com/api/v1/users",
            }
        ],
        "security": {"findings": {}},
    }

    normalized = scorer._normalize_results_payload(payload)

    assert normalized["metadata"]["discovery_baseline"][0]["path"] == "/api/v1/users"
    assert normalized["metadata"]["traffic_capture"]["captured_request_count"] == 1


def test_scorer_recovers_json_from_raw_output():
    scorer = CertificationScorer(
        {
            "usability": 0.1,
            "security": 0.2,
            "efficiency": 0.05,
            "reliability": 0.15,
            "performance": 0.15,
            "accessibility": 0.1,
            "scalability": 0.15,
            "stability": 0.1,
        }
    )

    raw_output = """
    CrewAI wrapped output:
    ```json
      {
        "metadata": {
          "requested_tool_count": 1,
          "resolved_tool_count": 1,
          "execution_state": "COMPLETED",
          "discovery_baseline": [
            {
              "method": "GET",
              "path": "/api/v1/api-docs"
            }
          ]
        },
        "security": {
          "findings": {
          "scenario_1_zap": {
            "success": true,
            "status": "PASS",
            "severity": "INFO",
            "mode": "LIVE",
            "path": "https://example.com/api",
            "log": "zap executed",
            "recommendation": "N/A"
          }
        }
      }
    }
    ```
    """

    normalized = scorer._normalize_results_payload({"raw_output": raw_output})

    assert normalized["metadata"]["resolved_tool_count"] == 1
    assert normalized["security"]["findings"]["scenario_1_zap"]["status"] == "PASS"
    assert normalized["metadata"]["discovery_baseline"][0]["path"] == "/api/v1/api-docs"


def test_excel_report_normalizes_target_paths(tmp_path):
    generator = ExcelReportGenerator(output_dir=str(tmp_path))
    raw_results = {
        "metadata": {"target_url": "https://www.myntra.com", "discovery_baseline": []},
        "security": {
            "findings": {
                "secure_config": {
                    "success": False,
                    "status": "FAIL",
                    "mode": "LIVE",
                    "severity": "HIGH",
                    "path": "https://www.w3schools.com/api/v1/health",
                    "log": "sample",
                    "recommendation": "fix",
                }
            }
        },
    }
    scorecard = {
        "usability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
        "security": {"passed": 0, "failed": 1, "skipped": 0, "total_tools": 1, "blockers": 1, "attribute_score": 0, "weight": 0.2, "weighted_contribution": 0, "verdict": "FAIL"},
        "efficiency": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.05, "weighted_contribution": 0, "verdict": "N/A"},
        "reliability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "performance": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "accessibility": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
        "scalability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "stability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
    }

    out_path = Path(generator.generate(raw_results, scorecard, 0.0, "Bronze"))
    wb = load_workbook(out_path, data_only=True)
    ws = wb["All Tool Findings"]

    assert ws["F2"].value == "https://www.myntra.com/api/v1/health"


def test_excel_report_uses_traffic_inventory_fallback(tmp_path):
    traffic_json = tmp_path / "traffic-inventory.json"
    traffic_json.write_text(
        json.dumps(
            {
                "capture_profile": {
                    "captured_requests": [
                        {
                            "method": "GET",
                            "url": "https://example.com/api/v1/users",
                            "path": "/api/v1/users",
                            "host": "example.com",
                            "status": 200,
                            "resource_type": "xhr",
                        }
                    ]
                },
                "captured_requests": [
                    {
                        "method": "GET",
                        "url": "https://example.com/api/v1/users",
                        "path": "/api/v1/users",
                        "host": "example.com",
                        "status": 200,
                        "resource_type": "xhr",
                    }
                ],
                "unique_api_count": 1,
            }
        ),
        encoding="utf-8",
    )

    generator = ExcelReportGenerator(output_dir=str(tmp_path))
    raw_results = {
        "metadata": {
            "target_url": "https://example.com",
            "discovery_baseline": [],
            "traffic_capture": {
                "traffic_inventory_json_path": str(traffic_json),
                "captured_requests": [],
            },
        },
        "security": {"findings": {}},
    }
    scorecard = {
        "usability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
        "security": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.2, "weighted_contribution": 0, "verdict": "N/A"},
        "efficiency": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.05, "weighted_contribution": 0, "verdict": "N/A"},
        "reliability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "performance": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "accessibility": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
        "scalability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "stability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
    }

    out_path = Path(generator.generate(raw_results, scorecard, 0.0, "Bronze", filepath="inventory_fallback.xlsx"))
    wb = load_workbook(out_path, data_only=True)
    ws = wb["Traffic Inventory"]

    assert ws.max_row >= 2
    assert ws["A2"].value == "GET"
    assert ws["C2"].value == "/api/v1/users"


def test_excel_report_uses_scenario_endpoint_fallback_for_baseline(tmp_path):
    generator = ExcelReportGenerator(output_dir=str(tmp_path))
    raw_results = {
        "metadata": {"target_url": "https://example.com", "discovery_baseline": []},
        "scenarios": [
            {
                "scenario_id": "scenario-1",
                "scenario_name": "Catalog",
                "endpoints": [
                    {
                        "method": "GET",
                        "url": "https://example.com/api/v1/products",
                        "source": "traffic-inventory",
                        "expected_behavior": "Returns products",
                    }
                ],
            }
        ],
        "security": {"findings": {}},
    }
    scorecard = {
        "usability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
        "security": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.2, "weighted_contribution": 0, "verdict": "N/A"},
        "efficiency": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.05, "weighted_contribution": 0, "verdict": "N/A"},
        "reliability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "performance": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "accessibility": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
        "scalability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.15, "weighted_contribution": 0, "verdict": "N/A"},
        "stability": {"passed": 0, "failed": 0, "skipped": 0, "total_tools": 0, "blockers": 0, "attribute_score": 0, "weight": 0.1, "weighted_contribution": 0, "verdict": "N/A"},
    }

    out_path = Path(generator.generate(raw_results, scorecard, 0.0, "Bronze", filepath="scenario_baseline_fallback.xlsx"))
    wb = load_workbook(out_path, data_only=True)
    ws = wb["Discovered API Baseline"]

    assert ws.max_row >= 2
    assert ws["A2"].value == "GET"
    assert ws["B2"].value == "/api/v1/products"


def test_discovery_inventory_filters_noise_and_keeps_api_requests(tmp_path):
    artifact_dir = tmp_path / "keploy" / "tests" / "keploy"
    artifact_dir.mkdir(parents=True)
    artifact_path = artifact_dir / "mixed.json"
    artifact_path.write_text(
        json.dumps(
            {
                "capture_profile": {
                    "captured_requests": [
                        {
                            "method": "GET",
                            "url": "https://example.com/",
                            "path": "/",
                            "host": "example.com",
                            "resource_type": "document",
                        },
                        {
                            "method": "GET",
                            "url": "https://example.com/style.css",
                            "path": "/style.css",
                            "host": "example.com",
                            "resource_type": "stylesheet",
                        },
                        {
                            "method": "GET",
                            "url": "https://example.com/api/v1/users",
                            "path": "/api/v1/users",
                            "host": "example.com",
                            "resource_type": "xhr",
                        },
                        {
                            "method": "POST",
                            "url": "https://example.com/auth/session",
                            "path": "/auth/session",
                            "host": "example.com",
                            "resource_type": "document",
                        },
                        {
                            "method": "GET",
                            "url": "https://cdn.example.net/asset.js",
                            "path": "/asset.js",
                            "host": "cdn.example.net",
                            "resource_type": "script",
                        },
                    ]
                }
            }
        ),
        encoding="utf-8",
    )

    inventory = build_discovery_inventory("https://example.com", workspace_root=tmp_path)
    discovered_paths = {entry["path"] for entry in inventory["discovered_apis"]}

    assert "/api/v1/users" in discovered_paths
    assert "/auth/session" in discovered_paths
    assert "/" not in discovered_paths
    assert "/style.css" not in discovered_paths
    assert inventory["inventory_summary"]["page_route_request_count"] >= 1
    assert inventory["inventory_summary"]["static_asset_request_count"] >= 1
